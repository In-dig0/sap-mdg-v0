"""
MDG — Migration Data Governance
Export Excel degli esiti dei controlli (condiviso da Cockpit e Check Results).
"""

import io
from datetime import date

import pandas as pd


def issues_filename(check_id: str, source_table: str) -> str:
    """Nome del file Excel: controllo, tabella di controllo e data (es. CK518-S_CUST_CONT#ZBP-AddInterlocutore_20261023.xlsx)."""
    return f"{check_id}-{source_table}_{date.today().strftime('%Y%m%d')}.xlsx"


def df_to_xlsx(
    df,
    sheet_name: str = "Risultati",
    legenda_rows: list | None = None,
) -> bytes:
    from openpyxl.styles import Font, PatternFill

    def _style_header(ws):
        from openpyxl.styles import Border, Side, Alignment
        header_font = Font(bold=True)
        header_fill = PatternFill("solid", start_color="BFBFBF", end_color="BFBFBF")
        border_side = Side(style="thin", color="000000")
        header_border = Border(
            left=border_side, right=border_side,
            top=border_side, bottom=border_side,
        )
        for cell in ws[1]:
            cell.font      = header_font
            cell.fill      = header_fill
            cell.border    = header_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        # Filtro automatico su tutto il range dati
        ws.auto_filter.ref = ws.dimensions

    def _autofit(ws):
        for col in ws.columns:
            max_len = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in col
            )
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 80)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        # Foglio principale
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws_main = writer.sheets[sheet_name]
        _style_header(ws_main)
        _autofit(ws_main)

        # Foglio Legenda
        df_legenda = pd.DataFrame(
            legenda_rows or [],
            columns=["Codice controllo", "Descrizione controllo", "Categoria", "Stato"],
        )
        df_legenda.to_excel(writer, index=False, sheet_name="Legenda")
        ws_leg = writer.sheets["Legenda"]
        _style_header(ws_leg)
        _autofit(ws_leg)

    return buf.getvalue()


def issues_to_xlsx(df_issues, check_id: str, description: str, category: str) -> bytes:
    """Excel dei record in Error/Warning di un controllo, con foglio Legenda."""
    legenda_rows = [
        {
            "Codice controllo":      check_id,
            "Descrizione controllo": description,
            "Categoria":             category,
            "Stato":                 stato,
        }
        for stato in sorted(df_issues["Stato"].unique())
    ]
    return df_to_xlsx(df_issues, sheet_name="Error+Warning", legenda_rows=legenda_rows)
